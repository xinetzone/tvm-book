
import logging
import numpy as np
import vta

logging.basicConfig(filename='build/test.log',
                    filemode="w",
                    format='%(asctime)s|%(levelname)s|%(name)s->%(funcName)s@%(message)s',
                    level=logging.DEBUG)


import numpy as np
import tvm
import vta
from tvm.ir.module import IRModule
from tvm.script import tir as T


from dataclasses import dataclass, asdict, astuple
import numpy as np
import pandas as pd


@dataclass
class Config:
    input_size: int
    output_size: int
    ci: int # 输入通道个数
    co: int # 输出通道个数
    kernel: int # 卷积核大小
    padding: int # padding 大小
    stride: int # 跨步大小


config = Config(2, 2, 16, 16, 3, 1, 1)


co_list = np.array([1<<k for k in range(6)])
h_list = np.array([1<<k for k in range(9)])
h_list.sort()
wls = set()
kernel = 3
for stride in [1, 2]:
    for padding in [0, 1, 2]:
        if stride==2:
            padding = 1
        for kernel in [3, 1]:
            for hi in h_list:
                ho = round(np.floor((hi - kernel + 2*padding)/stride) + 1)
                if stride==1 and hi!=ho:
                    continue
                for co in co_list:
                    if stride==2:
                        ci = co//2
                    else:
                        ci = co
                    if ho > hi:
                        continue
                    if ci == 0:
                        ci = 1
                    config = (hi, ho, ci*16, co*16, kernel, padding, stride)
                    wls.add(config)
                    logging.info(f"workload: {config}")


from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

names = []
for inp_name in [128, 256, 384, 512]:
    for acc_name in [128, 256, 384, 512]:
        name = f"inp={inp_name},acc={acc_name}"
        names.append(name)
wb = Workbook()
ws = wb.active
ws.title = "(inp,wgt=18,acc)SRAM 对比"
column_names = ["input_size", "output_size", "ci", "co", "kernel", "padding", "stride",
                *names]
ws.append(column_names)
# for col in "EFGHIJKLMNOPQRSTUVWXYZ":
#     ws.column_dimensions[col].width = 12
# ws.row_dimensions[0].height=20
# 文本自动换行
for r in ws:
    for c in r:
        c.alignment=Alignment(wrapText=True)
wls = sorted(wls)
for wl in wls:
    ws.append(wl)
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
tab = Table(displayName="卷积算子性能", ref=f"A1:W{len(wls)+1}")
tab.tableStyleInfo = style
ws.add_table(tab)
wb.save("build/常用卷积性能和sram相关性分析测试表.xlsx")








